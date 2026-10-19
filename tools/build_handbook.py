"""Generate the fictional Meridian Works handbook as real DOCX and PDF files.

Source text lives in handbook_src/*.md (clauses numbered like 4.2).
Output goes to data/handbook/. PDFs are made by converting DOCX with LibreOffice.
Run from the project root:  python tools/build_handbook.py
"""
import re
import subprocess
import sys
from pathlib import Path

import yaml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "handbook_src"
OUT = ROOT / "data" / "handbook"

# Policies delivered as PDF. Everything else stays DOCX.
PDF_IDS = {"POL-04", "POL-05", "POL-08", "POL-12", "POL-14", "POL-02_v3.0"}

# Real tables that restate thresholds, so the parser has to handle tables.
TABLES = {
    "POL-04": ("Approval matrix", ["Claim amount (RM)", "Approver"],
               [["Up to 2,000", "Direct manager"],
                ["Above 2,000 and up to 10,000", "Department head"],
                ["Above 10,000", "CFO"]]),
    "POL-05": ("Travel allowance schedule", ["Item", "Tier 1 city", "Tier 2 city", "Domestic"],
               [["Hotel cap per night (RM)", "450", "300", "Not applicable"],
                ["Meal allowance per day (RM)", "180", "130", "70"]]),
    "POL-08": ("Gift thresholds", ["Value of gift received (RM)", "Required action"],
               [["Up to 200", "May accept if occasional"],
                ["Above 200 and up to 500", "Declare to Compliance within 7 days"],
                ["Above 500", "Refuse or return"]]),
    "POL-12": ("Purchasing thresholds", ["Purchase value (RM)", "Requirement", "Approver"],
               [["Below 5,000", "One quotation", "Manager"],
                ["5,000 to 50,000", "Three written quotations", "Department head"],
                ["Above 50,000", "Formal tender", "CFO and CEO"]]),
}


def add_field(run, code):
    for kind, text in (("begin", None), (None, code), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def parse(path):
    text = path.read_text(encoding="utf-8")
    m = re.match(r"---\n(.*?)\n---\n", text, re.S)
    meta = yaml.safe_load(m.group(1))
    return meta, text[m.end():].strip().splitlines()


def shade(cell, color):
    tc = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tc.append(shd)


def build(path, meta, lines, out_docx):
    pid = meta["doc_id"]
    restricted = meta["access"] != ["all"]
    superseded = meta["status"] == "superseded"
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)

    sec = d.sections[0]
    hp = sec.header.paragraphs[0]
    banner = "Meridian Works Sdn Bhd | Internal"
    if restricted:
        banner = "Meridian Works Sdn Bhd | RESTRICTED: " + ", ".join(meta["access"]).upper()
    if superseded:
        banner = "Meridian Works Sdn Bhd | SUPERSEDED, DO NOT USE FOR CURRENT DECISIONS"
    hp.text = banner
    hp.runs[0].font.size = Pt(9)
    hp.runs[0].font.color.rgb = RGBColor(0xB0, 0x20, 0x20) if (restricted or superseded) else RGBColor(0x55, 0x55, 0x55)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run(f"{pid} v{meta['version']}  |  Page ")
    r.font.size = Pt(9)
    r2 = fp.add_run()
    r2.font.size = Pt(9)
    add_field(r2, "PAGE")

    title = [l for l in lines if l.startswith("# ")][0][2:]
    d.add_heading(title, level=1)

    ctl = [("Document ID", pid), ("Version", str(meta["version"])),
           ("Effective date", str(meta["effective_date"])), ("Status", meta["status"].capitalize()),
           ("Policy owner", meta["owner"]),
           ("Classification", "Restricted" if restricted else "Internal"),
           ("Audience", ", ".join(meta["access"]))]
    if meta.get("supersedes"):
        ctl.append(("Supersedes", meta["supersedes"]))
    if meta.get("superseded_by"):
        ctl.append(("Superseded by", meta["superseded_by"]))
    t = d.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for k, v in ctl:
        row = t.add_row().cells
        row[0].text, row[1].text = k, v
        shade(row[0], "E8EEF7")
        row[0].paragraphs[0].runs[0].bold = True
    d.add_paragraph()

    table_done = False
    for line in lines:
        if line.startswith("# "):
            continue
        if line.startswith("## "):
            d.add_heading(line[3:], level=2)
            continue
        m = re.match(r"^(\d+\.\d+) (.*)$", line)
        if m:
            p = d.add_paragraph()
            p.paragraph_format.left_indent = Pt(24)
            p.paragraph_format.first_line_indent = Pt(-24)
            b = p.add_run(m.group(1) + "  ")
            b.bold = True
            p.add_run(m.group(2))
        elif line.strip():
            d.add_paragraph(line)

    key = pid if not superseded else f"{pid}_v{meta['version']}"
    if pid in TABLES and not superseded and not table_done:
        cap, head, rows = TABLES[pid]
        d.add_heading(f"Schedule: {cap}", level=2)
        tb = d.add_table(rows=1, cols=len(head))
        tb.style = "Table Grid"
        for i, h in enumerate(head):
            tb.rows[0].cells[i].text = h
            shade(tb.rows[0].cells[i], "D9E2F3")
            tb.rows[0].cells[i].paragraphs[0].runs[0].bold = True
        for row in rows:
            cells = tb.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = v
    d.save(out_docx)
    return key


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    pdf_jobs = []
    for f in sorted(SRC.glob("*.md")):
        meta, lines = parse(f)
        pid = meta["doc_id"]
        suffix = f"_v{meta['version']}"
        stem = f.stem.replace(".md", "")
        stem = re.sub(r"_v3\.\d.*$", "", stem) + (suffix if pid == "POL-02" else "")
        docx_path = OUT / f"{stem}.docx"
        key = build(f, meta, lines, docx_path)
        if key in PDF_IDS:
            pdf_jobs.append(docx_path)
        print("built", docx_path.name)
    for p in pdf_jobs:
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(OUT), str(p)],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        p.unlink()
        print("converted to PDF", p.stem)


if __name__ == "__main__":
    sys.exit(main())
